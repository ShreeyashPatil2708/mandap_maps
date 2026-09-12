#!/usr/bin/env python3
"""Generate backend/db/seed-data.json from the two source spreadsheets.

Reads:
  data/MandapMaps_Ganpati_Data_2026.xlsx  (main record per Ganpati)
  data/MandapMaps_Metro_Food_2026.xlsx    (metro + food companion, joined by ID)

Writes:
  backend/db/seed-data.json  (one object per Ganpati, matching the seed columns)

Requires openpyxl:  pip install openpyxl
The spreadsheets and the generated JSON are the private dataset and are
gitignored; this script (code only) is committed.
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MAIN_XLSX = ROOT / "data" / "MandapMaps_Ganpati_Data_2026.xlsx"
COMPANION_XLSX = ROOT / "data" / "MandapMaps_Metro_Food_2026.xlsx"
OUT = ROOT / "backend" / "db" / "seed-data.json"


def clean(v):
    """Trim strings; treat blanks and 'TO ADD'/'TO UPDATE' placeholders as None."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if s.upper().startswith(("TO ADD", "TO UPDATE", "TBD", "N/A")):
        return None
    return s


def to_num(v):
    s = clean(v)
    if s is None:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def leading_int(v):
    """First integer in a string, e.g. '2 - Famous' -> 2, '3rd Manacha' -> 3."""
    s = clean(v)
    if s is None:
        return None
    m = re.search(r"\d+", s)
    return int(m.group()) if m else None


def rows_of(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return header, rows[1:]


def col(index_map, prefix):
    """Column index whose header starts with `prefix` (avoids pasting the
    spreadsheet's em-dash-laden header text into this file)."""
    for name, i in index_map.items():
        if name.startswith(prefix):
            return i
    raise KeyError(prefix)


EM_DASH = "\u2014"  # spreadsheet uses U+2014 as the food-cell field separator


def parse_metro(station, walk):
    """'Kasba Peth (Purple Line)' + '~5 min walk' -> {name, line, dist}."""
    station = clean(station)
    if station is None:
        return None
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", station)
    if m:
        name, line = m.group(1).strip(), m.group(2).strip()
    else:
        name, line = station, None
    return {"name": name, "line": line, "dist": clean(walk)}


def split_on_emdash(s):
    return re.split(r"\s*" + re.escape(EM_DASH) + r"\s*", s)


def parse_food(cell):
    """'Name , desc , distance' separated by em dashes -> {name, type, dist}."""
    s = clean(cell)
    if s is None:
        return None
    parts = [p.strip() for p in split_on_emdash(s) if p.strip()]
    if not parts:
        return None
    if len(parts) >= 3:
        return {"name": parts[0], "type": " ".join(parts[1:-1]), "dist": parts[-1]}
    if len(parts) == 2:
        return {"name": parts[0], "type": parts[1], "dist": None}
    return {"name": parts[0], "type": None, "dist": None}


CORRECTIONS_FILE = Path(__file__).with_name("coordinate-corrections.json")


def load_corrections():
    """
    Hand-checked coordinate fixes, keyed by name. Each one records the value it
    replaces, so it applies only while the spreadsheet still holds that stale
    value: once a coordinate is corrected at the source, the entry goes quiet
    instead of overwriting the newer value. See the notes in the JSON file.
    """
    if not CORRECTIONS_FILE.exists():
        return {}
    doc = json.loads(CORRECTIONS_FILE.read_text(encoding="utf-8"))
    return {c["name_english"]: c for c in doc.get("corrections", [])}


def corrected(name, lat, lng, corrections):
    """The coordinate to seed: the correction when the sheet is still stale."""
    fix = corrections.get(name)
    if not fix or lat is None or lng is None:
        return lat, lng
    was_lat, was_lng = fix["from"]
    if round(lat, 6) == round(was_lat, 6) and round(lng, 6) == round(was_lng, 6):
        return fix["to"][0], fix["to"][1]
    return lat, lng


def maps_url(sheet_value, lat, lng):
    """
    A map link that actually resolves.

    The companion sheet's "Google Maps Link" column holds invented short links
    (maps.app.goo.gl/<PandalName>), which 404: a real short link is a random
    code, never a readable name, and there is no way to tell a good one from a
    made-up one by looking at it. So short links are dropped, a full
    google.com/maps URL is trusted, and anything else falls back to the
    coordinates, which is exactly what the app's "Open in Google Maps" button
    builds. Without coordinates there is no link, rather than a broken one.
    """
    value = (sheet_value or "").strip()
    if "google.com/maps" in value and "goo.gl" not in value:
        return value
    if lat is None or lng is None:
        return None
    return f"https://www.google.com/maps/search/?api=1&query={lat},{lng}"


def main():
    m_hdr, m_rows = rows_of(MAIN_XLSX, "Pune Ganpati 2026")
    c_hdr, c_rows = rows_of(COMPANION_XLSX, "MandapMaps Companion")

    mi = {name: i for i, name in enumerate(m_hdr)}
    ci = {name: i for i, name in enumerate(c_hdr)}
    companion = {clean(r[ci["ID"]]): r for r in c_rows}

    corrections = load_corrections()

    records = []
    for r in m_rows:
        rid = clean(r[mi["ID"]])
        comp = companion.get(rid)

        tier = leading_int(r[mi["Tier"]])
        is_manache_tier = "manache" in str(r[mi["Tier"]] or "").lower()
        rank = leading_int(r[mi["Manacha / Rank"]])
        manacha_number = rank if (is_manache_tier and rank and 1 <= rank <= 5) else None

        tags_raw = clean(r[mi["Tags (app filters)"]]) or ""
        tags = [t.strip() for t in tags_raw.split("|") if t.strip()]

        metro, food, gmaps = [], [], None
        if comp is not None:
            for s, w in (("Nearest Metro Station 1 (Line)", "Approx Walk from Metro 1"),
                         ("Nearest Metro Station 2 (Line)", "Approx Walk from Metro 2")):
                mm = parse_metro(comp[ci[s]], comp[ci[w]])
                if mm:
                    metro.append(mm)
            for prefix in ("Food Spot 1", "Food Spot 2", "Food Spot 3"):
                ff = parse_food(comp[col(ci, prefix)])
                if ff:
                    food.append(ff)
            gmaps = clean(comp[ci["Google Maps Link (Pandal / Temple)"]])

        name_english = clean(r[mi["Name (English)"]])
        lat, lng = corrected(
            name_english, to_num(r[mi["Latitude"]]), to_num(r[mi["Longitude"]]), corrections
        )
        gmaps = maps_url(gmaps, lat, lng)

        records.append({
            "name_english": name_english,
            "name_marathi": clean(r[mi["Name (Marathi)"]]),
            "manacha_number": manacha_number,
            "tier": tier,
            "category": clean(r[mi["Category"]]),
            "area": clean(r[mi["Area / Neighbourhood"]]),
            "year_established": clean(r[mi["Year Established"]]),
            "history_english": clean(r[mi["History (English)"]]),
            "history_marathi": clean(r[mi["History (Marathi)"]]),
            "significance_short": clean(r[col(mi, "Why Significant")]),
            "idol_description": clean(r[mi["Idol / Murti Description"]]),
            "mandir_address": clean(r[mi["Mandir Address (permanent)"]]),
            "pandal_address": clean(r[col(mi, "Pandal Address")]),
            "latitude": lat,
            "longitude": lng,
            "morning_aarti": clean(r[mi["Morning Aarti Time"]]),
            "evening_aarti": clean(r[mi["Evening Aarti Time"]]),
            "special_events": clean(r[mi["Special Events During 10 Days"]]),
            "tags": tags,
            "photo_url": clean(r[mi["Photo URL"]]),
            "google_maps_url": gmaps,
            "metro": metro,
            "food": food,
            "is_manacha": manacha_number is not None,
        })

    OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(records)} records to {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
